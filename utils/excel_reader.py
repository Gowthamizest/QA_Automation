import os
from typing import List, Dict, Optional
import openpyxl


class ExcelReader:
    """Simple Excel reader for .xlsx files."""

    def __init__(self, file_path: str):
        if not file_path:
            raise ValueError("Excel file path cannot be empty")

        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Excel file not found: {file_path}")

        self.file_path = file_path
        self.workbook = openpyxl.load_workbook(file_path, data_only=True)

    def get_sheet_names(self) -> List[str]:
        return self.workbook.sheetnames

    def read_sheet(self, sheet_name: str, header_row: int = 1) -> List[Dict[str, object]]:
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {self.file_path}")

        sheet = self.workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            return []

        if header_row < 1 or header_row > len(rows):
            raise IndexError("header_row is out of range")

        headers = [str(cell).strip() if cell is not None else "" for cell in rows[header_row - 1]]
        nonblank = [h.casefold() for h in headers if h]
        if len(nonblank) != len(set(nonblank)):
            raise ValueError(f"Duplicate column headers in worksheet '{sheet_name}'.")
        if not any(headers):
            raise ValueError(f"Worksheet '{sheet_name}' has no column headers in row {header_row}.")
        data_rows: List[Dict[str, object]] = []

        for row in rows[header_row:]:
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue

            record: Dict[str, object] = {}
            for index, header in enumerate(headers):
                value = row[index] if index < len(row) else None
                record[header or f"column_{index + 1}"] = value
            data_rows.append(record)

        return data_rows

    def read_rows(self, sheet_name: str, header_row: int = 1) -> List[List[object]]:
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Sheet '{sheet_name}' not found in {self.file_path}")

        sheet = self.workbook[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))

        if not rows:
            return []

        if header_row < 1 or header_row > len(rows):
            raise IndexError("header_row is out of range")

        return [list(row) for row in rows[header_row:]]

    def read_row(self, sheet_name: str, row_number: int, header_row: int = 1) -> Optional[Dict[str, object]]:
        rows = self.read_sheet(sheet_name, header_row=header_row)
        if 0 < row_number <= len(rows):
            return rows[row_number - 1]
        return None

    def close(self) -> None:
        self.workbook.close()
